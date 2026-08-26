#!/usr/bin/env python3
from __future__ import annotations

import re
import subprocess
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT / ".local" / "external" / "oh-my-code"
SOURCE = ROOT / "oh_my_code_agent_harness_textbook.md"
OUTPUT = Path.home() / "Desktop" / "Oh_My_Code_Agent_Harness_从零到生产_源码教学与面试手册.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F2F2F2"
DARK_GRAY = "404040"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, east_asia: str = "等线", ascii_font: str = "Aptos") -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_keep_with_next(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "在 Word 中右键此处并选择“更新域”以刷新目录。"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2, placeholder, fld_char3])


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "等线", "Menlo")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(80, 80, 80)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run)
            run.bold = True
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in (
        ("Title", 28, BLUE),
        ("Heading 1", 19, BLUE),
        ("Heading 2", 15, DARK_GRAY),
        ("Heading 3", 12, DARK_GRAY),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in doc.styles:
        style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Code Block"]
    style.font.name = "Menlo"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    style.font.size = Pt(8.5)
    style.paragraph_format.left_indent = Cm(0.45)
    style.paragraph_format.right_indent = Cm(0.25)
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(6)
    p_pr = style._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)

    if "Quote" not in doc.styles:
        quote = doc.styles.add_style("Quote", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = doc.styles["Quote"]
    quote.font.italic = False
    quote.font.color.rgb = RGBColor(55, 55, 55)
    quote.paragraph_format.left_indent = Cm(0.6)
    quote.paragraph_format.right_indent = Cm(0.3)
    quote.paragraph_format.space_before = Pt(4)
    quote.paragraph_format.space_after = Pt(6)
    q_pr = quote._element.get_or_add_pPr()
    q_shd = OxmlElement("w:shd")
    q_shd.set(qn("w:fill"), LIGHT_BLUE)
    q_pr.append(q_shd)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("从一次 Tool Call 到稳定的 Agent Harness")
    set_run_font(run, "等线", "Aptos Display")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("oh-my-code 全量源码教学与大模型应用面试手册")
    set_run_font(run)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)

    doc.add_paragraph()
    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_runs(scope, "覆盖 Agent 架构、DAG/Gate、Graph Truth、证据新鲜度、并发控制、恢复、多宿主适配、MCP、Language Runtime 与工程测试")

    for _ in range(5):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_runs(meta, "源码基线：2a984718（2026-08-22）\n生成日期：2026-08-25\n内部敏感信息已脱敏")
    doc.add_page_break()

    p = doc.add_paragraph("阅读说明", style="Heading 1")
    set_paragraph_keep_with_next(p)
    doc.add_paragraph(
        "本手册以本地源码为事实基线，并吸收相关文章的设计动机。文章中的早期描述若与当前代码冲突，以当前代码、测试和正式架构文档为准。"
    )
    doc.add_paragraph(
        "正文按学习路径组织；附录按源码路径组织。即使以后无法访问代码，也可以依靠正文理解机制，并用附录恢复模块级细节。"
    )
    doc.add_paragraph("目录", style="Heading 1")
    add_toc(doc.add_paragraph())
    doc.add_page_break()


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int] | None:
    if index + 1 >= len(lines):
        return None
    header = lines[index].strip()
    separator = lines[index + 1].strip()
    if not (header.startswith("|") and separator.startswith("|")):
        return None
    if not re.match(r"^\|?\s*:?-{3,}", separator):
        return None
    rows: list[list[str]] = []
    cursor = index
    while cursor < len(lines) and lines[cursor].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[cursor].strip().strip("|").split("|")]
        if cursor != index + 1:
            rows.append(cells)
        cursor += 1
    return rows, cursor


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index in range(cols):
            value = values[col_index] if col_index < len(values) else ""
            cells[col_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[col_index])
            if row_index == 0:
                set_cell_shading(cells[col_index], BLUE)
            paragraph = cells[col_index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
        if row_index == 0:
            set_repeat_table_header(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def append_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                p.paragraph_format.keep_together = True
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, "等线", "Menlo")
                in_code = False
                code_lines = []
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        parsed = parse_table(lines, index)
        if parsed:
            rows, index = parsed
            add_table(doc, rows)
            continue
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("# "):
            heading = stripped[2:]
            if heading.startswith("第") or heading in {"导读：这本手册解决什么问题", "结语：真正稳定的 Agent 不是更会“想”，而是更难错误地“完成”", "附录说明"}:
                if len(doc.paragraphs) > 3:
                    doc.add_page_break()
            doc.add_paragraph(heading, style="Heading 1")
        elif stripped.startswith("## "):
            doc.add_paragraph(stripped[3:], style="Heading 2")
        elif stripped.startswith("### "):
            doc.add_paragraph(stripped[4:], style="Heading 3")
        elif stripped.startswith("> "):
            p = doc.add_paragraph(style="Quote")
            add_inline_runs(p, stripped[2:])
        elif re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, text)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, stripped)
        index += 1


def git_files() -> list[Path]:
    result = subprocess.run(
        ["git", "-C", str(REPO), "ls-files", "-z"],
        check=True,
        stdout=subprocess.PIPE,
    )
    return [REPO / item.decode("utf-8") for item in result.stdout.split(b"\0") if item]


def line_count(path: Path) -> str:
    try:
        data = path.read_bytes()
    except OSError:
        return "?"
    if b"\0" in data:
        return "binary"
    return str(data.count(b"\n") + (0 if not data or data.endswith(b"\n") else 1))


def file_role(relative: str) -> str:
    name = Path(relative).name
    suffix = Path(relative).suffix
    rules = [
        ("src/lib/core/harness", "Harness 状态、控制、证据或生命周期核心"),
        ("src/lib/core/", "运行时共享基础设施"),
        ("src/lib/language/lsp/", "LSP 连接、会话、文档或协议实现"),
        ("src/lib/language/providers/", "语言能力 Provider"),
        ("src/lib/language/router/", "语言请求路由与策略"),
        ("src/lib/language/facade/", "Language MCP 对外工具契约"),
        ("src/lib/language/ast", "AST 搜索/替换"),
        ("src/lib/language/", "语言智能运行时"),
        ("src/hooks/targets/", "宿主 Hook 协议适配"),
        ("src/hooks/", "Hook 事件处理与上下文注入"),
        ("src/integrations/targets/", "宿主安装、卸载、状态与资产同步"),
        ("src/integrations/", "Target Provider 和宿主注册"),
        ("src/mcp/", "MCP Server 入口"),
        ("src/cli/", "CLI 参数与命令"),
        ("src/build/", "声明式资产编译与生成"),
        ("src/opencode-plugin/", "OpenCode 插件和事件桥"),
        ("src/dsh-", "DSH 插件集成"),
        ("src/lib/acp/", "ACP Agent/Session 客户端"),
        ("agents/", "声明式 Agent 角色契约"),
        ("skills/", "声明式 Skill/工作流契约"),
        ("orchestration/common/", "宿主无关正确性与 Assurance 契约"),
        ("orchestration/hosts/", "宿主专用 Orchestrator Bundle"),
        ("targets/", "目标宿主静态配置与 Manifest"),
        ("test/unit/", "单元/契约测试"),
        ("test/smoke/", "构建产物和宿主最小链路 Smoke"),
        ("test/e2e/", "安装与真实流程 E2E"),
        ("test/fixtures/", "测试夹具"),
        ("test/", "测试基础设施或静态验证器"),
        ("docs/design/architecture/", "当前架构契约"),
        ("docs/design/proposals/", "设计提案，未必已落地"),
        ("docs/design/deprecated/", "历史设计，仅供演进参考"),
        ("docs/research/", "外部方案研究与比较"),
        ("docs/", "产品、开发或使用文档"),
        ("deepwiki/", "生成的源码导航 Wiki"),
        ("scripts/", "构建、检查、Smoke 或发布脚本"),
        (".github/", "GitHub 协作模板"),
        (".codebase/", "CI 配置"),
        ("drafts/", "未进入正式能力面的草稿"),
        ("dockers/", "容器化测试环境"),
    ]
    for prefix, role in rules:
        if relative.startswith(prefix):
            return role
    if name == "package.json":
        return "包元数据、依赖和脚本"
    if name.startswith("tsconfig"):
        return "TypeScript 编译配置"
    if name.startswith("vitest"):
        return "Vitest 测试配置"
    if name == "README.md":
        return "项目入口说明"
    if suffix in {".md"}:
        return "项目文档"
    if suffix in {".json", ".yaml", ".yml"}:
        return "声明式配置"
    return "项目支撑文件"


def first_signal(path: Path) -> str:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return ""
    for pattern in (
        r"^description:\s*[|>]?\s*(.+)$",
        r"^#\s+(.+)$",
        r"^export\s+(?:async\s+)?(?:function|class|const|type|interface)\s+([A-Za-z0-9_]+)",
    ):
        match = re.search(pattern, text, re.MULTILINE)
        if match:
            value = re.sub(r"\s+", " ", match.group(1)).strip()
            return value[:90]
    return ""


def append_source_inventory(doc: Document, files: list[Path]) -> None:
    doc.add_page_break()
    doc.add_paragraph("附录 A：全量源码文件索引", style="Heading 1")
    doc.add_paragraph(
        f"本附录覆盖 canonical 工作树的全部 {len(files)} 个 Git tracked 文件。职责说明由路径、文件标题、声明和导出符号共同生成；它用于定位，不代表每个提案均已落地。"
    )

    groups: dict[str, list[Path]] = {}
    for path in files:
        relative = path.relative_to(REPO).as_posix()
        top = relative.split("/", 1)[0] if "/" in relative else "根目录"
        groups.setdefault(top, []).append(path)

    for group in sorted(groups):
        doc.add_paragraph(group, style="Heading 2")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["路径", "行数", "模块职责", "首要信号"]
        for i, value in enumerate(headers):
            set_cell_shading(table.rows[0].cells[i], BLUE)
            p = table.rows[0].cells[i].paragraphs[0]
            r = p.add_run(value)
            set_run_font(r)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8)
        set_repeat_table_header(table.rows[0])
        for path in sorted(groups[group]):
            relative = path.relative_to(REPO).as_posix()
            row = table.add_row().cells
            values = [relative, line_count(path), file_role(relative), first_signal(path)]
            for i, value in enumerate(values):
                set_cell_margins(row[i], 45, 55, 45, 55)
                p = row[i].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(value)
                set_run_font(r, "等线", "Menlo" if i == 0 else "Aptos")
                r.font.size = Pt(7)


def append_module_map(doc: Document, files: list[Path]) -> None:
    doc.add_page_break()
    doc.add_paragraph("附录 B：模块规模与阅读顺序", style="Heading 1")
    categories = [
        ("Harness Core", "src/lib/core/harness"),
        ("Core Other", "src/lib/core/"),
        ("Language", "src/lib/language/"),
        ("Hooks", "src/hooks/"),
        ("Integrations", "src/integrations/"),
        ("MCP", "src/mcp/"),
        ("OpenCode Plugin", "src/opencode-plugin/"),
        ("CLI/Build", "src/cli/"),
        ("Agents", "agents/"),
        ("Skills", "skills/"),
        ("Orchestration", "orchestration/"),
        ("Targets", "targets/"),
        ("Tests", "test/"),
        ("Docs", "docs/"),
        ("DeepWiki", "deepwiki/"),
    ]
    rows = [["模块", "文件数", "建议阅读重点"]]
    for label, prefix in categories:
        matches = [p for p in files if p.relative_to(REPO).as_posix().startswith(prefix)]
        rows.append([label, str(len(matches)), file_role(prefix)])
    add_table(doc, rows)
    doc.add_paragraph(
        "推荐源码顺序：README/overview -> orchestrator Skill -> Agent contracts -> harness types/operations -> harness core -> completion/proof/attempt/snapshot -> hooks/session -> target provider -> concrete integrations -> language runtime -> tests。"
    )


def append_test_map(doc: Document, files: list[Path]) -> None:
    doc.add_page_break()
    doc.add_paragraph("附录 C：测试文件与行为覆盖索引", style="Heading 1")
    test_files = [p for p in files if p.relative_to(REPO).as_posix().startswith("test/")]
    rows = [["测试文件", "测试块", "主要覆盖"]]
    for path in sorted(test_files):
        relative = path.relative_to(REPO).as_posix()
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            text = ""
        count = len(re.findall(r"\b(?:it|test)\s*\(", text))
        rows.append([relative, str(count), file_role(relative)])
    add_table(doc, rows)


def append_glossary(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("附录 D：核心术语速查", style="Heading 1")
    rows = [
        ["术语", "定义"],
        ["Acceptance Horizon", "最终必须交付或证明的完整结果边界。"],
        ["Commitment Graph", "当前信息下已确定必要并可安全派发的最小任务图。"],
        ["Graph Truth", "持久化的节点、依赖、Gate、控制和恢复事实源。"],
        ["Work State", "节点局部源码、日志、工具输出和临时推理。"],
        ["Gate", "阻止错误传播的验收边界，不是普通 Todo。"],
        ["Assurance", "Gate 对完成的约束强度：advisory、gated、ralph。"],
        ["Node", "一个语义工作或验收边界。"],
        ["Attempt", "某节点的一次具体派发和执行。"],
        ["Reconcile", "根据新观察重算未派发图后缀。"],
        ["Frozen Prefix", "已派发/完成、不可被后续计划改写的历史。"],
        ["Proof Freshness", "证据仍绑定当前合同、快照和相关历史的性质。"],
        ["Snapshot Binding", "代码/工作树内容状态的稳定标识。"],
        ["Resource Claim", "Attempt 对路径或命名资源声明的访问范围。"],
        ["Fencing Token", "阻止旧 Attempt 或错误 Owner 提交状态的令牌。"],
        ["Q6 Boundary", "材料性终态观察完成后进行权威收敛决策的边界。"],
        ["Q6-Live", "只消费 Sideband 事件的受限中途控制，不产生验收事实。"],
        ["Session Binding", "宿主 session/target/cwd 与 Harness 的精确关联。"],
        ["Stop Closeout", "宿主想停止时执行完成检查、延续或有界释放。"],
        ["Target Adapter", "隔离宿主安装、Hook、MCP 和生命周期差异。"],
    ]
    add_table(doc, rows)


def append_source_notes(doc: Document, files: list[Path]) -> None:
    doc.add_page_break()
    doc.add_paragraph("附录 E：资料来源、版本差异与脱敏说明", style="Heading 1")
    paragraphs = [
        f"源码事实源：本地 `.local/external/oh-my-code`，提交 `2a98471805f357cb3fc0da68d1b59225ab2dd9c8`。共索引 {len(files)} 个 tracked 文件。",
        "重复工作树：`.local/external/oh-my-code-ssh` 与 canonical 工作树提交和文件内容一致，因此未重复计入。",
        "文章来源：ByteTech 文章《如何设计一个面向复杂需求和长任务的 Agent Harness 框架》。手册吸收其问题定义、Graph State/Work State、角色分离、任务图、Gate 和 Graph Truth 思想，但没有逐字转载。",
        "事实优先级：当前源码与测试 > current architecture 文档 > README/DeepWiki > 文章 > proposals/research/deprecated。提案只作为演进讨论，不写成已落地能力。",
        "脱敏：未保留作者身份、工号、邮箱、内部仓库地址、内部安装源、文档 Token、消息 ID 或认证信息。命令示例使用通用占位和公开式表达。",
        "覆盖方法：逐文件建立路径、行数、类型、职责和首要声明索引；逐模块检查核心导出、入口、配置和测试；对 Harness、Proof、Attempt、Snapshot、Session、Hook、Target、Language 等关键链路进行源码级交叉阅读。",
    ]
    for text in paragraphs:
        doc.add_paragraph(text)


def add_headers_and_footers(doc: Document) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("oh-my-code Agent Harness 源码教学与面试手册")
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 120, 120)
        add_page_number(section.footer.paragraphs[0])


def enable_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def validate_docx(path: Path, expected_files: int) -> None:
    if not path.exists() or path.stat().st_size < 100_000:
        raise RuntimeError("DOCX output is missing or unexpectedly small")
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names:
            raise RuntimeError("Invalid DOCX: word/document.xml missing")
        xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        required = [
            "从一次 Tool Call 到稳定的 Agent Harness",
            "Acceptance Horizon",
            "harness-completion-policy.ts",
            "附录 A：全量源码文件索引",
            f"全部 {expected_files} 个 Git tracked 文件",
        ]
        for marker in required:
            if marker not in xml:
                raise RuntimeError(f"DOCX validation marker missing: {marker}")


def main() -> None:
    files = git_files()
    markdown = SOURCE.read_text(encoding="utf-8")

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    append_markdown(doc, markdown)
    append_source_inventory(doc, files)
    append_module_map(doc, files)
    append_test_map(doc, files)
    append_glossary(doc)
    append_source_notes(doc, files)
    add_headers_and_footers(doc)
    enable_update_fields(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "从一次 Tool Call 到稳定的 Agent Harness"
    doc.core_properties.subject = "oh-my-code 全量源码教学与大模型应用面试手册"
    doc.core_properties.author = "源码研读整理（脱敏版）"
    doc.core_properties.keywords = "Agent Harness, DAG, Gate, MCP, LLM, long-running agent"
    doc.save(OUTPUT)
    validate_docx(OUTPUT, len(files))
    print(OUTPUT)
    print(f"files={len(files)} size={OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
